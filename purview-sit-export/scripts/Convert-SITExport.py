"""Build a readable Excel workbook from a Purview custom SIT export folder.

Usage: python Convert-SITExport.py "<export folder>" [-o output.xlsx]
"""
import argparse
import csv
import json
import math
import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

MCE = "http://schemas.microsoft.com/office/2011/mce"
NS = {"m": MCE}
GUID_RE = re.compile(r"^[0-9a-f]{8}-([0-9a-f]{4}-){3}[0-9a-f]{12}$", re.IGNORECASE)

CONF_FILL = {
    "High": PatternFill("solid", fgColor="1F4E79"),
    "Medium": PatternFill("solid", fgColor="2E75B6"),
    "Low": PatternFill("solid", fgColor="BDD7EE"),
}
CONF_FONT = {
    "High": Font(bold=True, color="FFFFFF"),
    "Medium": Font(bold=True, color="FFFFFF"),
    "Low": Font(bold=True, color="1F3864"),
}
BAND_FILLS = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3F8")]
SIT_BORDER = Border(top=Side(style="medium", color="1F3864"))
PATTERN_BORDER = Border(top=Side(style="thin", color="A6A6A6"))
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TABLE_STYLE = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
# Keyed by Get-DlpKeywordDictionary Identity, which is the GUID SITs reference
DICTIONARIES_BY_ID = {}


def q(tag):
    return f"{{{MCE}}}{tag}"


def local(tag):
    return tag.rsplit("}", 1)[-1]


def confidence_label(value):
    level = int(value)
    if level >= 85:
        return "High"
    if level >= 75:
        return "Medium"
    return "Low"


def proximity_label(value):
    if not value:
        return ""
    if str(value).lower() == "unlimited":
        return "Anywhere in document"
    return f"{value} characters"


def load_xml(path):
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = data.decode("utf-16")
    elif data[:2] == b"<\x00":
        text = data.decode("utf-16-le")
    elif data[:2] == b"\x00<":
        text = data.decode("utf-16-be")
    else:
        text = data.decode("utf-8-sig")
    # ElementTree rejects str input that still declares a UTF-16 encoding
    text = re.sub(r"^\s*<\?xml[^>]*\?>", "", text)
    return ET.fromstring(text)


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def resource_text(root, sit_id, tag):
    resource = root.find(f".//m:LocalizedStrings/m:Resource[@idRef='{sit_id}']", NS)
    if resource is None:
        return ""
    node = resource.find(f"m:{tag}[@default='true']", NS)
    if node is None:
        node = resource.find(f"m:{tag}", NS)
    return (node.text or "").strip() if node is not None else ""


def build_catalog(root):
    catalog = {}
    for regex in root.iter(q("Regex")):
        catalog[regex.get("id")] = {
            "type": "Regular expression",
            "definition": regex.text or "",
            "match_style": "",
            "case": "",
        }
    for keyword in root.iter(q("Keyword")):
        terms, styles, cases = [], set(), set()
        for group in keyword.findall("m:Group", NS):
            styles.add(group.get("matchStyle", "word").capitalize())
            for term in group.findall("m:Term", NS):
                terms.append(term.text or "")
                cases.add("Yes" if term.get("caseSensitive", "false").lower() == "true" else "No")
        catalog[keyword.get("id")] = {
            "type": "Keyword list",
            "definition": "\n".join(terms),
            "match_style": ", ".join(sorted(styles)),
            "case": "Mixed" if len(cases) > 1 else next(iter(cases), ""),
        }
    return catalog


def resolve(element_id, catalog, sit_names):
    if element_id in catalog:
        return catalog[element_id]
    if element_id.lower() in DICTIONARIES_BY_ID:
        entry = DICTIONARIES_BY_ID[element_id.lower()]
        terms = [t.strip() for t in (entry.get("Terms") or "").splitlines() if t.strip()]
        return {"type": "Keyword dictionary", "definition": f"[{entry.get('Name', '')}]\n" + "\n".join(terms),
                "match_style": (entry.get("MatchStyle") or "").capitalize(), "case": ""}
    if element_id.lower() in sit_names:
        return {"type": "Sensitive info type", "definition": sit_names[element_id.lower()], "match_style": "", "case": ""}
    if element_id.startswith("Func_"):
        return {"type": "Built-in function", "definition": element_id, "match_style": "", "case": ""}
    if GUID_RE.match(element_id):
        return {
            "type": "Keyword dictionary",
            "definition": "Referenced by ID only - see the Keyword Dictionaries sheet for dictionary terms.",
            "match_style": "",
            "case": "",
        }
    return {"type": "Other", "definition": "", "match_style": "", "case": ""}


def group_rule(any_node):
    low, high = any_node.get("minMatches"), any_node.get("maxMatches")
    if low is None and high is None:
        return "at least 1 of"
    if high is None:
        return f"at least {low} of"
    if low is None:
        return f"up to {high} of"
    if low == high:
        return f"exactly {low} of"
    return f"{low} to {high} of"


def make_element(role, requirement, match_node, catalog, sit_names):
    element_id = match_node.get("idRef")
    if match_node.get("proximity"):
        requirement += f" (within {proximity_label(match_node.get('proximity'))})"
    return {"role": role, "requirement": requirement, "element": element_id, **resolve(element_id, catalog, sit_names)}


def add_any_group(any_node, label, elements, catalog, sit_names):
    requirement = f"{label} - {group_rule(any_node)}"
    nested = 0
    for child in any_node:
        if local(child.tag) == "Match":
            elements.append(make_element("Supporting", requirement, child, catalog, sit_names))
        elif local(child.tag) == "Any":
            nested += 1
            add_any_group(child, f"{label}.{nested}", elements, catalog, sit_names)


def parse_entity_pattern(pattern, catalog, sit_names):
    elements = []
    group_no = 0
    for child in pattern:
        tag = local(child.tag)
        if tag == "IdMatch":
            elements.append(make_element("Primary", "Required", child, catalog, sit_names))
        elif tag == "Match":
            elements.append(make_element("Supporting", "Required", child, catalog, sit_names))
        elif tag == "Any":
            group_no += 1
            add_any_group(child, f"Group {group_no}", elements, catalog, sit_names)
    elements.sort(key=lambda e: e["role"] != "Primary")
    return elements


def parse_rule_package(root, summary_by_id, sit_names, rule_pack_id):
    catalog = build_catalog(root)
    sits = []

    for entity in root.iter(q("Entity")):
        sit_id = entity.get("id").lower()
        if sit_id not in summary_by_id:
            continue
        patterns = [
            {
                "confidence": int(pattern.get("confidenceLevel")),
                "proximity": proximity_label(pattern.get("proximity") or entity.get("patternsProximity")),
                "elements": parse_entity_pattern(pattern, catalog, sit_names),
            }
            for pattern in entity.findall("m:Pattern", NS)
        ]
        modified = entity.get("lastModifiedTime", "")
        sits.append({
            "id": sit_id,
            "type": "Pattern-based",
            "recommended": entity.get("recommendedConfidence"),
            "default_proximity": proximity_label(entity.get("patternsProximity")),
            "last_modified": f"{modified[:16].replace('T', ' ')} UTC" if modified else "",
            "patterns": patterns,
        })

    fingerprints = {fp.get("id"): fp for fp in root.iter(q("Fingerprint"))}
    processors = {af.get("id"): af for af in root.iter(q("AdvancedFingerprint"))}
    for affinity in root.iter(q("Affinity")):
        sit_id = affinity.get("id").lower()
        if sit_id not in summary_by_id:
            continue
        patterns = []
        for evidence in affinity.findall("m:Evidence", NS):
            for match in evidence.findall("m:Match", NS):
                fingerprint = fingerprints.get(match.get("idRef"))
                if fingerprint is None:
                    continue
                processor = processors.get(fingerprint.get("linkedProcessorIdRef"))
                configs = processor.findall(".//m:FingerprintConfig", NS) if processor is not None else []
                levels = [(c.get("confidenceLevel"), c.get("threshold")) for c in configs] or [
                    (evidence.get("confidenceLevel"), fingerprint.get("threshold"))
                ]
                for confidence, threshold in levels:
                    patterns.append({
                        "confidence": int(confidence),
                        "proximity": "n/a",
                        "elements": [{
                            "role": "Primary",
                            "requirement": "Required",
                            "element": fingerprint.get("id"),
                            "type": "Document fingerprint",
                            "definition": f"Document is at least {threshold}% similar to the fingerprinted template document.",
                            "match_style": "",
                            "case": "",
                        }],
                    })
        sits.append({
            "id": sit_id,
            "type": "Document fingerprint",
            "recommended": affinity.get("thresholdConfidenceLevel"),
            "default_proximity": "n/a",
            "last_modified": "",
            "patterns": patterns,
        })

    for sit in sits:
        summary = summary_by_id[sit["id"]]
        sit["name"] = resource_text(root, sit["id"], "Name") or summary.get("Name", "")
        sit["description"] = resource_text(root, sit["id"], "Description") or summary.get("Description", "")
        sit["publisher"] = summary.get("Publisher", "")
        sit["rule_pack"] = rule_pack_id
    return sits


def write(cell, value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def estimate_height(values, widths):
    lines = 1
    for value, width in zip(values, widths):
        if value in (None, ""):
            continue
        chars = max(width - 2, 1)
        count = sum(max(1, math.ceil(len(line) / chars)) for line in str(value).split("\n"))
        lines = max(lines, count)
    return min(409, 15 * lines)


def write_table(ws, start_row, headers, widths, rows, name):
    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        write(ws.cell(row=start_row, column=col), header)
        ws.column_dimensions[get_column_letter(col)].width = width
    for offset, values in enumerate(rows, start=1):
        row = start_row + offset
        for col, value in enumerate(values, start=1):
            write(ws.cell(row=row, column=col), value).alignment = WRAP_TOP
        ws.row_dimensions[row].height = estimate_height(values, widths)
    last_row = start_row + max(len(rows), 1)
    table = Table(displayName=name, ref=f"A{start_row}:{get_column_letter(len(headers))}{last_row}")
    table.tableStyleInfo = TABLE_STYLE
    ws.add_table(table)


def build_workbook(sits, dictionaries, export_info, output):
    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    patterns_ws = wb.create_sheet("Patterns")
    dictionaries_ws = wb.create_sheet("Keyword Dictionaries")
    notes_ws = wb.create_sheet("Notes")

    # Patterns: one row per element, grouped by SIT then pattern in rule package order (matches the portal)
    headers = ["SIT Name", "Pattern", "Confidence", "Pattern Proximity", "Role", "Requirement",
               "Element", "Element Type", "Definition", "Match Style", "Case Sensitive"]
    widths = [34, 11, 12, 20, 12, 28, 34, 20, 70, 12, 10]
    rows, row_meta, sit_first_row = [], [], {}
    for sit in sits:
        sit_first_row[sit["id"]] = len(rows) + 2
        for number, pattern in enumerate(sit["patterns"], start=1):
            label = confidence_label(pattern["confidence"])
            for index, element in enumerate(pattern["elements"]):
                rows.append([sit["name"], f"Pattern {number}", label, pattern["proximity"], element["role"],
                             element["requirement"], element["element"], element["type"],
                             element["definition"], element["match_style"], element["case"]])
                row_meta.append((sit["id"], number == 1 and index == 0, index == 0, label))
    write_table(patterns_ws, 1, headers, widths, rows, "tblPatterns")

    band, previous_sit = -1, None
    for offset, (sit_id, sit_start, pattern_start, label) in enumerate(row_meta, start=2):
        if sit_id != previous_sit:
            band, previous_sit = band + 1, sit_id
        for col in range(1, len(headers) + 1):
            cell = patterns_ws.cell(row=offset, column=col)
            cell.fill = BAND_FILLS[band % 2]
            if sit_start:
                cell.border = SIT_BORDER
            elif pattern_start:
                cell.border = PATTERN_BORDER
        if sit_start:
            patterns_ws.cell(row=offset, column=1).font = Font(bold=True)
        confidence_cell = patterns_ws.cell(row=offset, column=3)
        confidence_cell.fill = CONF_FILL[label]
        confidence_cell.font = CONF_FONT[label]
        confidence_cell.alignment = Alignment(horizontal="center", vertical="top")
    patterns_ws.freeze_panes = "B2"

    # Overview
    title = "Microsoft Purview - Custom Sensitive Information Types"
    if export_info.get("Tenant"):
        title += f" - {export_info['Tenant']}"
    write(overview["A1"], title).font = Font(bold=True, size=14)
    details = []
    if export_info.get("Tenant"):
        tenant = export_info["Tenant"]
        details.append(f"Tenant: {tenant} ({export_info['TenantId']})" if export_info.get("TenantId") else f"Tenant: {tenant}")
    exported = f"Exported {export_info['ExportedAt']}"
    if export_info.get("ExportedBy"):
        exported += f" by {export_info['ExportedBy']}"
    details.append(exported)
    write(overview["A2"], " | ".join(details))
    write(overview["A3"], "Configuration snapshot. Select a SIT name to jump to its patterns.")
    for ref in ("A2", "A3"):
        overview[ref].font = Font(italic=True, color="595959")
    headers = ["SIT Name", "Type", "Description", "Publisher", "Recommended Confidence", "Default Proximity",
               "Patterns", "High", "Medium", "Low", "Last Modified", "SIT ID", "Rule Package ID"]
    widths = [42, 20, 60, 24, 15, 18, 10, 8, 10, 8, 20, 38, 38]
    rows = []
    for sit in sits:
        counts = [sum(confidence_label(p["confidence"]) == level for p in sit["patterns"]) for level in ("High", "Medium", "Low")]
        recommended = confidence_label(sit["recommended"]) if sit["recommended"] else ""
        rows.append([sit["name"], sit["type"], sit["description"], sit["publisher"], recommended,
                     sit["default_proximity"], len(sit["patterns"]), *counts, sit["last_modified"],
                     sit["id"], sit["rule_pack"]])
    write_table(overview, 4, headers, widths, rows, "tblOverview")
    for offset, sit in enumerate(sits, start=5):
        link = overview.cell(row=offset, column=1)
        link.hyperlink = f"#'Patterns'!A{sit_first_row[sit['id']]}"
        link.font = Font(color="0563C1", underline="single")
        recommended = overview.cell(row=offset, column=5)
        if recommended.value in CONF_FILL:
            recommended.fill = CONF_FILL[recommended.value]
            recommended.font = CONF_FONT[recommended.value]
    overview.freeze_panes = "B5"

    # Keyword dictionaries
    rows = []
    for entry in sorted(dictionaries, key=lambda d: d.get("Name", "").lower()):
        terms = [t.strip() for t in (entry.get("Terms") or "").splitlines() if t.strip()]
        rows.append([entry.get("Name", ""), entry.get("Description", ""), (entry.get("MatchStyle") or "").capitalize(),
                     len(terms), "\n".join(terms), entry.get("Identity", "")])
    write_table(dictionaries_ws, 1, ["Dictionary Name", "Description", "Match Style", "Term Count", "Terms", "Dictionary ID"],
                [55, 40, 12, 12, 60, 38], rows, "tblDictionaries")
    dictionaries_ws.freeze_panes = "A2"

    # Notes / legend
    notes = [
        ("Confidence levels", "Purview stores confidence as a number. High = 85, Medium = 75, Low = 65 "
                              "(values of 85+ show as High, 75-84 as Medium, below 75 as Low)."),
        ("Pattern", "A SIT can have several patterns. A match on any one pattern triggers the SIT at that "
                    "pattern's confidence level. Pattern numbers follow the order shown in the Purview portal; "
                    "filter the Confidence column to view a single level."),
        ("Role - Primary", "The element that must be found for the pattern to match."),
        ("Role - Supporting", "Additional evidence that must be found within the pattern proximity window."),
        ("Requirement", "'Required' means the element must be present. 'Group N - at least X of' means X "
                        "elements from that group must be present."),
        ("Pattern Proximity", "Character window around the primary element in which supporting elements must appear."),
        ("Keyword dictionary", "Dictionaries are referenced by ID inside the SIT definition. Their names and terms "
                               "are listed on the Keyword Dictionaries sheet."),
        ("Document fingerprint", "Only the similarity thresholds are shown. The original template document is not "
                                 "stored in or recoverable from the export."),
        ("Source", "Generated from the rule package XML, SIT summary and keyword dictionary exports from "
                   "Security & Compliance PowerShell."),
    ]
    notes_ws.column_dimensions["A"].width = 24
    notes_ws.column_dimensions["B"].width = 100
    write(notes_ws["A1"], "Notes").font = Font(bold=True, size=14)
    for row, (term, text) in enumerate(notes, start=3):
        write(notes_ws.cell(row=row, column=1), term).font = Font(bold=True)
        write(notes_ws.cell(row=row, column=2), text).alignment = WRAP_TOP
        notes_ws.row_dimensions[row].height = estimate_height([text], [100])

    wb.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("folder", type=Path, help="Folder containing SIT-Summary.csv and RulePack-*.xml")
    parser.add_argument("-o", "--output", type=Path, help="Output .xlsx (default: <folder>\\SIT-Configuration.xlsx)")
    args = parser.parse_args()

    summary_path = args.folder / "SIT-Summary.csv"
    if not summary_path.exists():
        sys.exit(f"SIT-Summary.csv not found in {args.folder}")

    summary = read_csv(summary_path)
    summary_by_id = {row["Id"].lower(): row for row in summary}
    sit_names = {row["Id"].lower(): row["Name"] for row in summary}

    dictionaries_path = args.folder / "KeywordDictionaries.csv"
    dictionaries = read_csv(dictionaries_path) if dictionaries_path.exists() else []
    DICTIONARIES_BY_ID.update({d["Identity"].strip().lower(): d for d in dictionaries if d.get("Identity")})

    sits = []
    for xml_path in sorted(args.folder.glob("RulePack-*.xml")):
        rule_pack_id = xml_path.stem.removeprefix("RulePack-")
        sits.extend(parse_rule_package(load_xml(xml_path), summary_by_id, sit_names, rule_pack_id))
    sits.sort(key=lambda s: s["name"].lower())

    info_path = args.folder / "ExportInfo.json"
    export_info = json.loads(info_path.read_text(encoding="utf-8-sig")) if info_path.exists() else {}
    if export_info.get("ExportedAtUtc"):
        stamp = datetime.strptime(export_info["ExportedAtUtc"], "%Y-%m-%dT%H:%M:%SZ")
        export_info["ExportedAt"] = stamp.strftime("%d %B %Y %H:%M UTC")
    else:
        export_info["ExportedAt"] = datetime.fromtimestamp(summary_path.stat().st_mtime).strftime("%d %B %Y")

    output = args.output or args.folder / "SIT-Configuration.xlsx"
    build_workbook(sits, dictionaries, export_info, output)
    print(f"{len(sits)} SITs, {sum(len(s['patterns']) for s in sits)} patterns -> {output}")


if __name__ == "__main__":
    main()
